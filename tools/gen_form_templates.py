#!/usr/bin/env python3
"""Sinh ba tệp mẫu C# từ reference/Mẫu.xlsx. CHẠY MỘT LẦN: từ đó sửa tay các tệp *Template.cs,
test TemplateMatchesMauTests giữ cho chúng đúng với Mẫu.xlsx.

    python3 -I gen_form_templates.py <Mẫu.xlsx> <thư mục đích>
"""
import re, sys, os, unicodedata
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

DOTS = re.compile(r"[.…]+")
DEFAULT_COL = 8.43
DEFAULT_ROW = 15.0

def esc(s):  # C# regular string literal (luôn NFC: Mẫu có một ô viết ở dạng NFD)
    s = unicodedata.normalize("NFC", s)
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "") + '"'

def addr(r, c): return f"{L(c)}{r}"

# ---- chú giải: ô -> danh sách đoạn. Mỗi đoạn là một tuple (loại, ...). -------------------------
# ("F", text)                      Fixed
# ("I", key, placeholder[, default[, hint]])  Input
# ("N", key)                       InputNumber
# ("S", key, placeholder)          SettingInput (Cài đặt)
# ("C", key)                       Computed
# ("O", key[, hint])               Overridable
# ("DOTS", [keys], strip)          tách theo dải chấm: chữ giữa các dải là Fixed, mỗi dải là Input
# ("TAIL", prefix, key, strip)     Fixed(prefix) + Input(key, phần còn lại của chữ Mẫu làm placeholder)
# ("HINT", prefix, key)            Fixed(prefix) + Input(key, "", hint = phần còn lại)
YEAR = "{năm}"
# khoá Computed mang CHỮ (không định dạng số) — Mẫu vẫn để định dạng #,##0 trên các ô này
TEXT_KEYS = {"PackageName", "RequestedInWords", "ContractorName", "BankAccount", "TotalInWords",
             "PaidInWords", "Line.Text"}

ANNOT = {
 "Thông tri": {
  "A1": [("S", "don_vi_cap_tren", "Đơn vị cấp trên")],
  "A2": [("S", "don_vi_co_so", "Đơn vị cơ sở")],
  "G2": [("DOTS", ["SoPhieu"], True)],
  "A3": [("DOTS", ["SoThongTri"], False)],
  "G3": [("TAIL", "Ghi sổ ngày", "NgayGhiSo", False)],
  "C8": [("DOTS", ["DonVi"], False)],
  "G8": [("DOTS", ["TaiKhoan1", "DonVi1", "SoChungTu1"], False)],
  "G9": [("DOTS", ["TaiKhoan2", "DonVi2", "SoChungTu2"], False)],
  "A9": [("TAIL", "Về", "VeThangNam", False)],
  "A10": [("DOTS", ["TenCTMT", "MaSo"], False)],
  "H24": [("C", "Total")],
  "C25": [("C", "TotalInWords")],
 },
 "KS chi": {
  "D3": [("TAIL", "Ngày", "NgayNhanGiayDeNghi", False)],
  "A4": [("DOTS", ["DonVi"], False)],
  "C7": [("C", "Accepted.Total")], "D7": [("C", "Accepted.NSNN")], "E7": [("C", "Accepted.NgoaiHoi")],
  "A8": [("DOTS", ["Muc1", "TieuMuc1"], False)],
  "A9": [("DOTS", ["Muc2", "TieuMuc2"], False)],
  "A10": [("DOTS", ["Muc3", "TieuMuc3"], False)],
  "A11": [("DOTS", ["Muc4", "TieuMuc4"], False)],
  "C13": [("C", "Recovered.Total")], "D13": [("C", "Recovered.NSNN")], "E13": [("C", "Recovered.NgoaiHoi")],
  "C14": [("N", "ThuHoiCacNamTruoc")], "C15": [("N", "ThuHoiNamNay")], "C16": [("N", "ThueGTGT")],
  "C17": [("C", "Paid.Total")], "D17": [("C", "Paid.NSNN")], "E17": [("C", "Paid.NgoaiHoi")],
  "A18": [("C", "PaidInWords")],
  "C19": [("N", "SoTuChoi")],
  "A20": [("TAIL", "Lý do từ chối: ", "LyDoTuChoi", False)],
  "B22": [("O", "GhiChu")],
 },
 "GĐNTT": {
  "A1": [("S", "don_vi_cap_tren", "Đơn vị cấp trên")],
  "A2": [("S", "don_vi_co_so", "Đơn vị cấp cơ sở")],
  "A3": [("TAIL", "Số: ", "So", False)],
  "D3": [("I", "NgayThangNam", "", "Hải phòng, ngày     tháng   năm " + YEAR, None)],
  "A6": [("HINT", "Kính gửi: ", "KinhGui")],
  "C8": [("O", "TenDuAnCongTrinh", "( Có thể gõ tự ra nội dung hoặc tự nhập liệu)")],
  "A9": [("HINT", " - Chủ đầu tư/Ban QLDA:  ", "ChuDauTu")],
  "A10": [("HINT", " - Căn cứ Quyết định số ", "CanCuQuyetDinh")],
  "A11": [("HINT", " - Căn cứ Hợp đồng số: ", "CanCuHopDong")],
  "A12": [("HINT", " - Căn cứ giấy đề nghị thanh toán vốn đầu tư số ", "CanCuGiayDeNghi")],
  "H13": [("O", "UnrecoveredAdvanceBefore", "( có thể gõ ra số liệu hoặc Tự nhập liệu)")],
  "C16": [("I", "ThuocKeHoach", "", None, None)],
  "A18": [("C", "PackageName")],
  "D18": [("C", "ContractTotal")], "F18": [("C", "CumulativeBefore")], "H18": [("C", "Requested")],
  "D19": [("C", "ContractTotal")], "F19": [("C", "CumulativeBefore")], "I19": [("C", "Requested")],
  "I20": [("C", "Requested")],
  "B21": [("C", "RequestedInWords")],
  "H23": [("C", "Recovered")],
  "H24": [("N", "ChuyenTienBaoHanh")],
  "H25": [("N", "ChuyenTienTamGiu")],
  "I26": [("C", "PaidToContractor")],
  "C27": [("C", "ContractorName")],
  "D28": [("C", "BankAccount")],
  "B29": [("I", "TaiNganHang", "", None, "Ngân hàng A ( Tự nhập)")],
 },
}

# ô thêm hoặc sửa so với Mẫu (chỗ lệch có chủ ý, ghi trong spec §5.8 của plan)
EXTRA_MERGES = {"GĐNTT": ["H13:I13", "H25:I25"]}
DROP_CELLS = {"GĐNTT": {"M13"}}
LAST_COL = {"Thông tri": 8, "KS chi": 5, "GĐNTT": 9}
LAST_ROW = {"Thông tri": 33, "KS chi": 34, "GĐNTT": 36}

# vùng lặp của Thông tri (hàng mẫu -> nguyên mẫu hàng dự án / vị trí / đợt / trống)
REGION = {"first": 15, "min": 9, "project": 15, "vitri": 16, "dot": 17, "blank": 20}
REGION_DEFAULTS = {"A": "010", "B": "011", "C": "9300", "D": "9301", "E": "00", "F": "00"}
REGION_KEYS = {"A": "LoaiNganSach", "B": "Khoan", "C": "Muc", "D": "TieuMuc", "E": "TietMuc", "F": "Nganh"}

HALIGN = {None: "General", "general": "General", "left": "Left", "center": "Center", "right": "Right", "justify": "Justify"}
VALIGN = {None: "Bottom", "bottom": "Bottom", "center": "Center", "top": "Top"}

def num(x):
    s = repr(float(x))
    if s.endswith(".0"): s = s[:-2]
    return s

def _side(cells, side):
    styles = {getattr(c.border, side).style for c in cells} - {None}
    if not styles: return None
    return "hair" if styles == {"hair"} else "thin"

def borders_str(cells):
    """hợp của các cạnh của vùng: chữ thường = nét thường, chữ HOA = nét hair (rất mảnh)."""
    rows = sorted({c.row for c in cells}); cols = sorted({c.column for c in cells})
    flags = ""
    for ch, side, edge in (("l", "left", [c for c in cells if c.column == cols[0]]),
                           ("r", "right", [c for c in cells if c.column == cols[-1]]),
                           ("t", "top", [c for c in cells if c.row == rows[0]]),
                           ("b", "bottom", [c for c in cells if c.row == rows[-1]])):
        st = _side(edge, side)
        if st == "thin": flags += ch
        elif st == "hair": flags += ch.upper()
    return flags

def style_expr(c, flags, is_money):
    f, a = c.font, c.alignment
    parts = [num(f.sz or 11)]
    kw = []
    if f.b: kw.append("bold: true")
    if f.i: kw.append("italic: true")
    if f.u: kw.append("underline: true")
    h = HALIGN[a.horizontal]
    if h != "General": kw.append(f"align: CellAlign.{h}")
    v = VALIGN[a.vertical]
    if v != "Bottom": kw.append(f"v: CellVAlign.{v}")
    if a.wrap_text: kw.append("wrap: true")
    if flags: kw.append(f'borders: "{flags}"')
    if is_money: kw.append(f"format: {esc(c.number_format)}")
    return "S(" + ", ".join(parts + kw) + ")"

def check_font(sheet, c):
    f = c.font
    if c.value is None: return
    if f.name != "Times New Roman":
        print(f"  CẢNH BÁO {sheet}!{c.coordinate}: font {f.name}", file=sys.stderr)
    if f.u or f.strike:
        print(f"  CẢNH BÁO {sheet}!{c.coordinate}: underline/strike", file=sys.stderr)

def segments(sheet, ws, c, text, strip_default=False):
    """trả về danh sách biểu thức C# cho các đoạn của ô (đã chú giải) hoặc Fixed(text)."""
    key = c.coordinate
    ann = ANNOT.get(sheet, {}).get(key)
    if ann is None:
        if isinstance(text, str) and text.startswith("="):
            raise SystemExit(f"{sheet}!{key}: công thức chưa chú giải: {text}")
        if text is None: return []
        if isinstance(text, (int, float)): text = str(text)
        text = text.replace("2026", YEAR)
        if DOTS.search(text) and len(DOTS.sub("", text)) != len(text) and "…" in text or ".." in text:
            print(f"  XEM LẠI {sheet}!{key}: có chấm chấm, chưa chú giải: {text!r}", file=sys.stderr)
        return [f"Tpl.Fixed({esc(text)})"]
    out = []
    for a in ann:
        kind = a[0]
        if kind == "F": out.append(f"Tpl.Fixed({esc(a[1])})")
        elif kind == "I":
            _, k, ph, *rest = a
            d = rest[0] if len(rest) > 0 else None
            hint = rest[1] if len(rest) > 1 else None
            out.append(f"Tpl.Input({esc(k)}, {esc(ph)}" + (f", @default: {esc(d)}" if d else "") + (f", hint: {esc(hint)}" if hint else "") + ")")
        elif kind == "N": out.append(f"Tpl.Number({esc(a[1])})")
        elif kind == "S": out.append(f"Tpl.Setting({esc(a[1])}, {esc(a[2])})")
        elif kind == "C": out.append(f"Tpl.Computed({esc(a[1])})")
        elif kind == "O": out.append(f"Tpl.Over({esc(a[1])}" + (f", hint: {esc(a[2])}" if len(a) > 2 else "") + ")")
        elif kind == "DOTS":
            keys, strip = a[1], a[2]
            t = str(text).replace("2026", YEAR)
            pieces, last = [], 0
            runs = list(DOTS.finditer(t))
            if len(runs) != len(keys):
                raise SystemExit(f"{sheet}!{key}: {len(runs)} dải chấm nhưng {len(keys)} khoá")
            for m, k in zip(runs, keys):
                pieces.append(("F", t[last:m.start()]))
                pieces.append(("I", k, m.group(0)))
                last = m.end()
            if t[last:]: pieces.append(("F", t[last:]))
            pieces = [p for p in pieces if not (p[0] == "F" and p[1] == "")]
            if strip:
                if pieces and pieces[0][0] == "F": pieces[0] = ("F", pieces[0][1].lstrip())
            for p in pieces:
                out.append(f"Tpl.Fixed({esc(p[1])})" if p[0] == "F" else f"Tpl.Input({esc(p[1])}, {esc(p[2])})")
        elif kind == "TAIL":
            prefix, k = a[1], a[2]
            t = str(text).replace("2026", YEAR)
            t = unicodedata.normalize("NFC", t)
            if not t.startswith(prefix): raise SystemExit(f"{sheet}!{key}: không bắt đầu bằng {prefix!r}: {t!r}")
            out.append(f"Tpl.Fixed({esc(prefix)})")
            out.append(f"Tpl.Input({esc(k)}, {esc(t[len(prefix):])})")
        elif kind == "HINT":
            prefix, k = a[1], a[2]
            t = str(text)
            if not t.startswith(prefix.rstrip()): raise SystemExit(f"{sheet}!{key}: không bắt đầu bằng {prefix!r}: {t!r}")
            rest = t[len(prefix.rstrip()):].strip()
            out.append(f"Tpl.Fixed({esc(prefix)})")
            out.append(f"Tpl.Input({esc(k)}, \"\", hint: {esc(rest)})")
    return out

def emit_sheet(wb, sheet, cls, out_path, region):
    ws = wb[sheet]
    last_col, last_row = LAST_COL[sheet], LAST_ROW[sheet]
    merges = {}
    covered = set()
    mranges = [str(r) for r in ws.merged_cells.ranges] + EXTRA_MERGES.get(sheet, [])
    for rng in mranges:
        a, b = rng.split(":")
        ca, cb = ws[a], ws[b]
        if ca.column > last_col: continue
        merges[(ca.row, ca.column)] = (cb.row - ca.row + 1, cb.column - ca.column + 1)
        for r in range(ca.row, cb.row + 1):
            for c in range(ca.column, cb.column + 1):
                if (r, c) != (ca.row, ca.column): covered.add((r, c))
    lines = []
    reg_rows = set()
    if region:
        reg_rows = set(range(region["first"], region["first"] + region["min"]))

    def cell_expr(row, col, relative=False, force=None):
        c = ws.cell(row, col)
        rs, cs = merges.get((row, col), (1, 1))
        group = [ws.cell(r, cc) for r in range(row, row + rs) for cc in range(col, col + cs)]
        flags = borders_str(group)
        text = c.value
        key = c.coordinate
        ann = force if force is not None else ANNOT.get(sheet, {}).get(key)
        is_money = bool(ann) and any(x[0] in ("C", "N", "O") and x[1] not in TEXT_KEYS and x[1] != "GhiChu"
                                     and x[1] != "TenDuAnCongTrinh" for x in ann) and "#,##0" in c.number_format
        check_font(sheet, c)
        st = style_expr(c, flags, is_money)
        if force is not None:
            ANNOT.setdefault(sheet, {})["__tmp__"] = force
            ANNOT[sheet][key] = force
        segs = segments(sheet, ws, c, text)
        if force is not None: del ANNOT[sheet][key]
        r = 0 if relative else row
        head = f"Tpl.Merged({r}, {col}, {rs}, {cs}" if (rs > 1 or cs > 1) else f"Tpl.Cell({r}, {col}"
        return f'{head}, {st}' + "".join(", " + s for s in segs) + ")"

    # ô thường
    body = []
    for row in range(1, last_row + 1):
        if row in reg_rows: continue
        for col in range(1, last_col + 1):
            if (row, col) in covered: continue
            if addr(row, col) in DROP_CELLS.get(sheet, set()): continue
            c = ws.cell(row, col)
            has_border = any(getattr(c.border, s).style for s in ("left", "right", "top", "bottom"))
            val = c.value
            if isinstance(val, str) and val.strip() == "": val = None
            is_anchor = (row, col) in merges
            if val is None and not has_border and not is_anchor: continue
            body.append("            " + cell_expr(row, col) + ",")

    # hàng mẫu của vùng lặp
    protos = {}
    if region:
        for name in ("project", "vitri", "dot", "blank"):
            r = region[name]
            h = ws.row_dimensions[r].height or DEFAULT_ROW
            cells = []
            for col in range(1, last_col + 1):
                letter = L(col)
                if name == "project" and col <= 6:
                    # ô Mục lục ngân sách: Input có mặc định
                    c = ws.cell(r, col)
                    flags = borders_str([c])
                    st = style_expr(c, flags, False)
                    cells.append(f'Tpl.Cell(0, {col}, {st}, Tpl.Input({esc(REGION_KEYS[letter])}, "", @default: {esc(REGION_DEFAULTS[letter])}))')
                elif name != "blank" and col == 7:
                    cells.append(cell_expr(r, col, relative=True, force=[("C", "Line.Text")]).replace("Tpl.Cell(0, 7,", "Tpl.Cell(0, 7,"))
                elif name != "blank" and col == 8:
                    c = ws.cell(r, col)
                    flags = borders_str([c])
                    st = style_expr(c, flags, True)
                    cells.append(f'Tpl.Cell(0, 8, {st}, Tpl.Computed("Line.Amount"))')
                else:
                    c = ws.cell(r, col)
                    flags = borders_str([c])
                    st = style_expr(c, flags, False)
                    cells.append(f"Tpl.Cell(0, {col}, {st})")
            protos[name] = (h, cells)

    heights = []
    for r in range(1, last_row + 1):
        if r in reg_rows: continue
        h = ws.row_dimensions[r].height
        if h: heights.append(f"[{r}] = {num(h)}")

    widths = []
    for col in range(1, last_col + 1):
        dim = None
        for k, d in ws.column_dimensions.items():
            if d.min and d.max and d.min <= col <= d.max: dim = d
        # openpyxl cho bề rộng LƯU TRONG TỆP (đã gồm 5px đệm). Excel hiện (và ClosedXML, và ExcelFitScale: px = w*7+5)
        # dùng bề rộng theo KÝ TỰ = (px − 5)/7 với px = làm tròn(lưu*7).
        if dim is not None and dim.width:
            w = round((round(dim.width * 7) - 5) / 7, 4)
        else:
            w = DEFAULT_COL
        widths.append(num(w))

    with open(out_path, "w", encoding="utf-8") as f:
        w = f.write
        w("// Sinh từ reference/Mẫu.xlsx bởi tools/gen_form_templates.py MỘT LẦN. Từ nay sửa tay tệp này;\n")
        w("// TemplateMatchesMauTests so nó với Mẫu.xlsx để không trôi khỏi mẫu.\n")
        w("using ConstructionFinance.Application.Printing;\n\n")
        w("namespace ConstructionFinance.Application.Forms;\n\n")
        w(f"/// Bản gốc của sheet \"{sheet}\" trong Mẫu.xlsx.\n")
        w(f"public static class {cls}\n{{\n")
        w("    private static CellStyle S(double size, bool bold = false, bool italic = false, bool underline = false,\n")
        w("        CellAlign align = CellAlign.General, CellVAlign v = CellVAlign.Bottom, bool wrap = false,\n")
        w("        string borders = \"\", string? format = null) =>\n")
        w("        Tpl.Style(size, bold, italic, underline, align, v, wrap, borders, format);\n\n")
        w("    public static FormTemplate Create() => new(\n")
        w(f"        SheetName: {esc(sheet)},\n")
        w("        Orientation: PrintOrientation.Portrait,\n")
        w("        ColumnWidths: new double[] { " + ", ".join(widths) + " },\n")
        w("        RowHeights: new Dictionary<int, double>\n        {\n")
        for i in range(0, len(heights), 6):
            w("            " + ", ".join(heights[i:i+6]) + ",\n")
        w("        },\n")
        w("        Cells: new[]\n        {\n")
        w("\n".join(body) + "\n")
        w("        },\n")
        if region:
            w("        Repeat: new RepeatRegion(\n")
            w(f"            FirstRow: {region['first']}, MinRows: {region['min']},\n")
            for i, name in enumerate(("project", "vitri", "dot", "blank")):
                h, cells = protos[name]
                label = {"project": "ProjectRow", "vitri": "ViTriRow", "dot": "DotRow", "blank": "BlankRow"}[name]
                w(f"            {label}: new TemplateRow({num(h)}, new[]\n            {{\n")
                for c in cells: w("                " + c + ",\n")
                w("            })" + (",\n" if i < 3 else ")") + "\n")
        else:
            w("        Repeat: null");
        w(");\n}\n")
    print("đã ghi", out_path, "—", len(body), "ô,", len(covered), "ô bị gộp che", file=sys.stderr)

if __name__ == "__main__":
    src, outdir = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    emit_sheet(wb, "Thông tri", "ThongTriTemplate", os.path.join(outdir, "ThongTriTemplate.cs"), REGION)
    emit_sheet(wb, "KS chi", "KsChiTemplate", os.path.join(outdir, "KsChiTemplate.cs"), None)
    emit_sheet(wb, "GĐNTT", "GdnttTemplate", os.path.join(outdir, "GdnttTemplate.cs"), None)
